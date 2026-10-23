import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.units import EMU_to_pixels

from bot_excel.core.models import AnnualSalesMonth, AnnualSalesRequest
from bot_excel.services.annual_consolidation import normalize_annual_months
from bot_excel.templates.registry import TemplateRegistry
from bot_excel.templates.ventas_anuales.generator import (
    AnnualSalesTemplateGenerator,
)


class AnnualSalesGenerationTests(unittest.TestCase):
    def test_generator_is_registered(self):
        self.assertIn(
            "ventas_consolidadas_anuales",
            TemplateRegistry().available(),
        )

    def test_rejects_invalid_year(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                AnnualSalesTemplateGenerator().generate(
                    AnnualSalesRequest(year=1999, output_dir=Path(tmp))
                )

    def test_rejects_invalid_or_duplicate_months(self):
        for months in (
            (AnnualSalesMonth(month=0),),
            (AnnualSalesMonth(month=13),),
            (AnnualSalesMonth(month=2), AnnualSalesMonth(month=2)),
        ):
            with self.subTest(months=months), self.assertRaises(ValueError):
                normalize_annual_months(months)

    def test_normalizes_to_twelve_months(self):
        normalized = normalize_annual_months((
            AnnualSalesMonth(month=3, real_sales=890),
        ))
        self.assertEqual(len(normalized), 12)
        self.assertEqual([item.month for item in normalized], list(range(1, 13)))
        self.assertEqual(normalized[2].real_sales, 890)
        self.assertIsNone(normalized[0].real_sales)

    def test_generates_annual_workbook_with_formulas_branding_and_charts(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = AnnualSalesTemplateGenerator().generate(
                AnnualSalesRequest(
                    year=2026,
                    output_dir=Path(tmp),
                    template_id="VA-2026-TEST",
                    months=(
                        AnnualSalesMonth(
                            month=1,
                            projected_sales=1_000_000,
                            real_sales=1_100_000,
                            projected_gc=1_000,
                            real_gc=950,
                            average_ticket=9_500,
                        ),
                        AnnualSalesMonth(
                            month=3,
                            projected_sales=0,
                            real_sales=50_000,
                            projected_gc=100,
                            real_gc=110,
                        ),
                        AnnualSalesMonth(
                            month=4,
                            projected_sales=900_000,
                            real_sales=750_000,
                            projected_gc=900,
                            real_gc=925,
                        ),
                        AnnualSalesMonth(
                            month=5,
                            projected_sales=800_000,
                            real_sales=880_000,
                            projected_gc=820,
                            real_gc=800,
                        ),
                    ),
                )
            )

            self.assertEqual(result.template_id, "VA-2026-TEST")
            self.assertEqual(result.period, "2026")
            self.assertEqual(
                result.workbook_path.name,
                "Ventas_Consolidadas_Anuales_2026_VA-2026-TEST.xlsx",
            )
            self.assertTrue(result.workbook_path.exists())
            self.assertTrue(result.qr_path.exists())

            workbook = load_workbook(result.workbook_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["CONSOLIDADO ANUAL", "Configuración"],
            )
            sheet = workbook["CONSOLIDADO ANUAL"]
            config = workbook["Configuración"]
            self.assertEqual(
                sheet["B1"].value,
                "VENTAS CONSOLIDADAS ANUALES 2026",
            )
            self.assertEqual(
                [sheet.cell(row=row, column=1).value for row in range(9, 21)],
                [
                    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
                    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE",
                    "NOVIEMBRE", "DICIEMBRE",
                ],
            )
            self.assertEqual(sheet["B9"].value, 1_000_000)
            self.assertEqual(sheet["C9"].value, 1_100_000)
            self.assertIn(sheet["B10"].value, (None, ""))
            self.assertEqual(sheet["B11"].value, 0)

            self.assertEqual(
                sheet["D9"].value,
                '=IF(OR(B9="",C9=""),"",C9-B9)',
            )
            self.assertEqual(sheet["E9"].value, '=IFERROR(C9/B9,"")')
            self.assertIn(sheet["F9"].value, (None, ""))
            self.assertEqual(sheet["F10"].value, '=IFERROR(C10/C9-1,"")')
            self.assertEqual(
                sheet["I9"].value,
                '=IF(OR(G9="",H9=""),"",H9-G9)',
            )
            self.assertEqual(sheet["K9"].value, "=SUM($B$9:B9)")
            self.assertEqual(sheet["K20"].value, "=SUM($B$9:B20)")
            self.assertEqual(sheet["L20"].value, "=SUM($C$9:C20)")
            self.assertEqual(sheet["M20"].value, '=IFERROR(L20/K20,"")')
            self.assertEqual(
                sheet["B21"].value,
                '=IF(COUNT(B9:B20)=0,"",SUM(B9:B20))',
            )
            self.assertEqual(sheet["E21"].value, '=IFERROR(C21/B21,"")')
            self.assertIn(sheet["J21"].value, (None, ""))
            self.assertEqual(sheet["M21"].value, '=IFERROR(L21/K21,"")')

            self.assertEqual(sheet["B3"].value, "=B21")
            self.assertIn("INDEX(A9:A20", sheet["B4"].value)
            self.assertEqual(sheet["H4"].value, "=G21")
            self.assertIn("AVERAGE(C9:C20)", sheet["E5"].value)
            self.assertEqual(len(sheet._charts), 2)
            self.assertEqual(
                [chart.title.tx.rich.p[0].r[0].t for chart in sheet._charts],
                [
                    "Ventas proyectadas vs reales",
                    "Cumplimiento mensual",
                ],
            )
            self.assertGreaterEqual(len(sheet.conditional_formatting), 2)

            visible_values = {
                cell.value
                for row in sheet.iter_rows()
                for cell in row
                if cell.value not in (None, "")
            }
            self.assertNotIn("VA-2026-TEST", visible_values)
            self.assertGreaterEqual(len(sheet._images), 2)
            qr_images = [
                image
                for image in sheet._images
                if image.anchor._from.col == 12
            ]
            self.assertEqual(len(qr_images), 1)
            self.assertEqual(EMU_to_pixels(qr_images[0].anchor.ext.cx), 48)
            self.assertEqual(
                config["B8"].value,
                "MCEXCEL|TIPO=VENTAS_ANUALES|"
                "CODIGO=VA-2026-TEST|PERIODO=2026",
            )
            self.assertEqual(config["B6"].value, 2026)
            self.assertEqual(config["B9"].value, "1.0")
            self.assertGreaterEqual(len(config._images), 1)

            self.assertEqual(
                sheet.print_area,
                "'CONSOLIDADO ANUAL'!$A$1:$M$41",
            )
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)
            self.assertTrue(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet["N9"].value, "=A9")
            self.assertEqual(sheet["O9"].value, '=IF(E9="","",E9)')
            self.assertEqual(sheet["P9"].value, 1)
            self.assertTrue(sheet.column_dimensions["N"].hidden)
            self.assertTrue(sheet.column_dimensions["O"].hidden)
            self.assertTrue(sheet.column_dimensions["P"].hidden)
            self.assertEqual(len(sheet._charts[1].series), 2)
            self.assertFalse(sheet._charts[1].visible_cells_only)

            formulas = [
                cell.value
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertFalse(any(
                error in formula
                for formula in formulas
                for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
            ))

            with zipfile.ZipFile(result.workbook_path) as archive:
                embedded = [
                    archive.read(name)
                    for name in archive.namelist()
                    if name.startswith("xl/media/")
                ]
                chart_parts = [
                    name
                    for name in archive.namelist()
                    if name.startswith("xl/charts/chart")
                ]
            self.assertGreaterEqual(len(embedded), 3)
            self.assertIn(result.qr_path.read_bytes(), embedded)
            self.assertEqual(len(chart_parts), 2)


if __name__ == "__main__":
    unittest.main()
