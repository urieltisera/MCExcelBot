import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from bot_excel.core.models import MonthlyTemplateRequest
from bot_excel.templates.ip.generator import IpTemplateGenerator


class IpHourDifferenceColorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        result = IpTemplateGenerator().generate(MonthlyTemplateRequest(
            year=2026,
            month=8,
            output_dir=Path(cls.temp_dir.name),
            projected_gc=(500, 500, 500),
            real_gc=(550, 450, 500),
            projected_hours=(100, 100, 100),
            real_hours=(120, 80, 100),
            ip_target=6.5,
            template_id="IP-2026-08-HCLR",
        ))
        cls.workbook = load_workbook(result.workbook_path, data_only=False)
        cls.sheet = cls.workbook["IP semanal y mensual"]
        cls.day_rows = {
            int(cell.value[-5:-3]): cell.row
            for cell in cls.sheet["A"]
            if isinstance(cell.value, str)
            and len(cell.value) >= 5
            and cell.value[-5:-3].isdigit()
        }

    @classmethod
    def tearDownClass(cls):
        cls.workbook.close()
        cls.temp_dir.cleanup()

    def _rules_for_cell(self, coordinate):
        cell = self.sheet[coordinate]
        for conditional_format in self.sheet.conditional_formatting:
            if any(
                cell.row in range(cell_range.min_row, cell_range.max_row + 1)
                and cell.column
                in range(cell_range.min_col, cell_range.max_col + 1)
                for cell_range in conditional_format.sqref.ranges
            ):
                return list(
                    self.sheet.conditional_formatting[conditional_format]
                )
        self.fail(f"No conditional-formatting rules found for {coordinate}")

    @staticmethod
    def _fill_rgb(rule):
        return rule.dxf.fill.fgColor.rgb[-6:]

    def test_hour_difference_formula_remains_real_minus_projected(self):
        expected_differences = {1: 20, 2: -20, 3: 0}
        for day, expected in expected_differences.items():
            row = self.day_rows[day]
            with self.subTest(day=day):
                self.assertEqual(
                    self.sheet[f"G{row}"].value,
                    f'=IF(OR(E{row}="",F{row}=""),"",F{row}-E{row})',
                )
                self.assertEqual(
                    self.sheet[f"F{row}"].value
                    - self.sheet[f"E{row}"].value,
                    expected,
                )

    def test_hour_difference_colors_are_positive_red_negative_green_zero_neutral(self):
        row = self.day_rows[1]
        rules = self._rules_for_cell(f"G{row}")
        by_comparison = {}
        for rule in rules:
            formula = rule.formula[0]
            if ">0" in formula:
                by_comparison["positive"] = self._fill_rgb(rule)
            elif "<0" in formula:
                by_comparison["negative"] = self._fill_rgb(rule)
            elif "=0" in formula:
                by_comparison["zero"] = self._fill_rgb(rule)

        self.assertEqual(by_comparison, {
            "positive": "FFC7CE",
            "negative": "C6EFCE",
            "zero": "FFE699",
        })

    def test_gc_and_ip_color_rules_are_unchanged(self):
        row = self.day_rows[1]
        gc_rules = self._rules_for_cell(f"D{row}")
        gc_by_comparison = {}
        for rule in gc_rules:
            formula = rule.formula[0]
            if ">0" in formula:
                gc_by_comparison["positive"] = self._fill_rgb(rule)
            elif "<0" in formula:
                gc_by_comparison["negative"] = self._fill_rgb(rule)
        self.assertEqual(gc_by_comparison, {
            "positive": "C6EFCE",
            "negative": "FFC7CE",
        })

        ip_rules = self._rules_for_cell(f"I{row}")
        self.assertEqual(len(ip_rules), 2)
        self.assertTrue(any(
            ">=H" in rule.formula[0]
            and self._fill_rgb(rule) == "C6EFCE"
            for rule in ip_rules
        ))
        self.assertTrue(any(
            "<H" in rule.formula[0]
            and self._fill_rgb(rule) == "FFC7CE"
            for rule in ip_rules
        ))


if __name__ == "__main__":
    unittest.main()
