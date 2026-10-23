import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import load_workbook
from openpyxl.utils.units import EMU_to_pixels

from bot_excel.branding import logo_path
from bot_excel.core.models import MonthlyTemplateRequest
from bot_excel.templates.ip.generator import IpTemplateGenerator
from bot_excel.templates.productividad_controlables.generator import (
    ProductivityControllablesTemplateGenerator,
)
from bot_excel.templates.registry import TemplateRegistry
from bot_excel.templates.ventas_gc.generator import SalesGcTemplateGenerator
class GenerationTests(unittest.TestCase):
    def assert_thin_gray_border(self, sheet, coordinates):
        for coordinate in coordinates:
            cell = sheet[coordinate]
            with self.subTest(sheet=sheet.title, cell=coordinate):
                for side in (
                    cell.border.left,
                    cell.border.right,
                    cell.border.top,
                    cell.border.bottom,
                ):
                    self.assertEqual(side.style, "thin")
                    self.assertEqual(side.color.type, "rgb")
                    self.assertEqual(side.color.rgb, "00808080")

    def body_merges(self, sheet, min_row, max_row):
        return [
            str(cell_range)
            for cell_range in sheet.merged_cells.ranges
            if cell_range.min_row <= max_row and cell_range.max_row >= min_row
        ]

    def assert_code_hidden_from_main_sheet(self, sheet, template_id):
        visible_values = {
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        }
        self.assertNotIn(template_id, visible_values)

    def test_generate_september(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = SalesGcTemplateGenerator().generate(MonthlyTemplateRequest(
                year=2026,
                month=9,
                output_dir=Path(tmp),
                projected_sales=(8_500_000, 9_000_000),
                projected_gc=(450, 450),
                template_id="VT-2026-09-TEST",
            ))
            self.assertTrue(result.workbook_path.exists())
            self.assertTrue(result.qr_path.exists())
            self.assertEqual(result.period, "2026-09")
            workbook = load_workbook(result.workbook_path, data_only=False)
            sheet = workbook["CONSOLIDADO"]
            self.assertEqual(
                sheet["B1"].value,
                "PROYECCION DE VENTAS SEPTIEMBRE 2026",
            )
            self.assert_code_hidden_from_main_sheet(sheet, "VT-2026-09-TEST")
            self.assertGreaterEqual(len(sheet._images), 2)
            qr_images = [
                image for image in sheet._images
                if image.anchor._from.col == 10
            ]
            self.assertEqual(len(qr_images), 1)
            self.assertEqual(EMU_to_pixels(qr_images[0].anchor.ext.cx), 48)
            self.assertEqual(sheet.column_dimensions["K"].width, 7)
            self.assertNotIn("L", sheet.column_dimensions)
            self.assertGreaterEqual(len(workbook["Configuración"]._images), 1)
            self.assertEqual(
                workbook["Configuración"]["B7"].value,
                "MCEXCEL|TIPO=VENTAS_GC|CODIGO=VT-2026-09-TEST|PERIODO=2026-09",
            )
            self.assertEqual(sheet["E4"].value, '=IF(OR(B4="",D4=""),"",D4-C4)')
            self.assertEqual(sheet["F4"].value, '=IF(E4="","",SUM($E$4:E4))')
            self.assertEqual(sheet["F5"].value, '=IF(E5="","",SUM($E$4:E5))')
            self.assertEqual(sheet["J5"].value, '=IF(I5="","",SUM($I$4:I5))')
            self.assertEqual(
                sheet["F20"].value,
                '=IF(E20="","",SUM($E$4:$E$18,$E$20:E20))',
            )
            self.assertEqual(
                sheet["J20"].value,
                '=IF(I20="","",SUM($I$4:$I$18,$I$20:I20))',
            )
            self.assertEqual(sheet["A19"].value, "1° QUINCENA")
            self.assertEqual(sheet["C19"].value, "=SUM(C4:C18)")
            self.assertEqual(sheet["C36"].value, "=SUM(C4:C18,C20:C35)")
            self.assertEqual(sheet["A3"].fill.fgColor.rgb, "00000000")
            self.assertEqual(sheet["A3"].font.color.rgb, "00FFFFFF")
            self.assertTrue(logo_path().is_file())
            self.assertIs(sheet.sheet_view.showGridLines, True)
            self.assertFalse(bool(sheet.print_options.gridLines))
            self.assertFalse(bool(sheet.print_options.gridLinesSet))
            self.assert_thin_gray_border(
                sheet,
                ("D5", "E5", "F5", "H5", "I5"),
            )
            self.assertCountEqual(
                self.body_merges(sheet, 3, 36),
                ["A19:B19", "A36:B36"],
            )
            self.assertIsNone(sheet["D35"].value)
            self.assert_thin_gray_border(sheet, ("D35",))
            self.assertGreaterEqual(sheet.column_dimensions["A"].width, 9)
            self.assertLessEqual(sheet.column_dimensions["A"].width, 10)
            daily_cells = [
                sheet.cell(row=row, column=1)
                for row in [*range(4, 19), *range(20, 35)]
            ]
            self.assertTrue(all(cell.value for cell in daily_cells))
            self.assertTrue(all(cell.font.color.rgb == "00FFFFFF" for cell in daily_cells))
            self.assertTrue(all(cell.fill.fgColor.rgb == "00000000" for cell in daily_cells))
            self.assertTrue(all(cell.border.left.style for cell in daily_cells))
            self.assertIn("MIÉRCOLES", [cell.value for cell in daily_cells])
            self.assertTrue(all(
                sheet.row_dimensions[cell.row].height <= 17
                for cell in daily_cells
            ))
            self.assertEqual(sheet.print_area, "'CONSOLIDADO'!$A$1:$K$36")
            with zipfile.ZipFile(result.workbook_path) as archive:
                embedded = [
                    archive.read(name)
                    for name in archive.namelist()
                    if name.startswith("xl/media/")
                ]
            self.assertGreaterEqual(len(embedded), 2)
            self.assertIn(result.qr_path.read_bytes(), embedded)
            result.qr_path.unlink()
            reopened = load_workbook(result.workbook_path)
            self.assertGreaterEqual(len(reopened["CONSOLIDADO"]._images), 2)

    def test_generate_productivity_and_controllables(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = ProductivityControllablesTemplateGenerator().generate(
                MonthlyTemplateRequest(
                    year=2026,
                    month=6,
                    output_dir=Path(tmp),
                    template_id="PC-2026-06-TEST",
                )
            )
            self.assertTrue(result.workbook_path.exists())
            self.assertTrue(result.qr_path.exists())
            self.assertEqual(result.period, "2026-06")

            workbook = load_workbook(result.workbook_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Productividad y controlables", "Configuración"],
            )
            sheet = workbook["Productividad y controlables"]
            self.assertEqual(
                sheet["B1"].value,
                "PROYECCION DE PRODUCTIVIDAD Y CONTROLABLES JUNIO 2026",
            )
            self.assert_code_hidden_from_main_sheet(
                sheet,
                "PC-2026-06-TEST",
            )
            self.assertGreaterEqual(len(sheet._images), 2)
            qr_images = [
                image for image in sheet._images
                if image.anchor._from.col == 16
            ]
            self.assertEqual(len(qr_images), 1)
            self.assertEqual(EMU_to_pixels(qr_images[0].anchor.ext.cx), 48)
            self.assertEqual(sheet.column_dimensions["Q"].width, 7)
            self.assertNotIn("R", sheet.column_dimensions)
            self.assertEqual(
                workbook["Configuración"]["B7"].value,
                "MCEXCEL|TIPO=PRODUCTIVIDAD_CONTROLABLES|"
                "CODIGO=PC-2026-06-TEST|PERIODO=2026-06",
            )
            self.assertEqual(sheet["E4"].value, '=IF(OR(B4="",D4=""),"",D4-C4)')
            self.assertEqual(sheet["F5"].value, '=IF(E5="","",F4+E5)')
            self.assertEqual(sheet["K4"].value, '=IF(OR(I4="",J4=""),"",I4-J4)')
            self.assertEqual(sheet["A19"].value, "1° QUINCENA")
            self.assertEqual(sheet["C19"].value, "=SUM(C4:C18)")
            self.assertEqual(
                sheet["G19"].value,
                '=IF(OR(D19="",D19=0),"",C19/D19)',
            )
            self.assertEqual(sheet["C36"].value, "=SUM(C4:C18,C20:C35)")
            self.assertEqual(sheet.print_area, "'Productividad y controlables'!$A$1:$Q$36")
            self.assertEqual(sheet["A3"].fill.fgColor.rgb, "00000000")
            self.assertEqual(sheet["A3"].font.color.rgb, "00FFFFFF")
            self.assertIs(sheet.sheet_view.showGridLines, True)
            self.assertFalse(bool(sheet.print_options.gridLines))
            self.assertFalse(bool(sheet.print_options.gridLinesSet))
            self.assert_thin_gray_border(
                sheet,
                (
                    "D5", "E5", "F5",
                    "J5",
                    "N5", "O5",
                ),
            )
            self.assertCountEqual(
                self.body_merges(sheet, 3, 36),
                ["A19:B19", "A36:B36"],
            )
            self.assertIsNone(sheet["D4"].value)
            self.assert_thin_gray_border(sheet, ("D4",))
            self.assertIsNone(sheet["J4"].value)
            self.assert_thin_gray_border(sheet, ("J4",))
            self.assertGreaterEqual(sheet.column_dimensions["A"].width, 9)
            self.assertLessEqual(sheet.column_dimensions["A"].width, 10)
            daily_cells = [
                sheet.cell(row=row, column=1)
                for row in [*range(4, 19), *range(20, 35)]
            ]
            self.assertTrue(all(cell.value for cell in daily_cells))
            self.assertTrue(all(cell.font.color.rgb == "00FFFFFF" for cell in daily_cells))
            self.assertTrue(all(cell.fill.fgColor.rgb == "00000000" for cell in daily_cells))
            self.assertTrue(all(cell.border.left.style for cell in daily_cells))
            self.assertIn("MIÉRCOLES", [cell.value for cell in daily_cells])
            self.assertTrue(all(
                sheet.row_dimensions[cell.row].height <= 16
                for cell in daily_cells
            ))

            formulas = [
                cell.value
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertEqual(len(formulas), 148)
            self.assertFalse(any(
                error in formula
                for formula in formulas
                for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
            ))
            with zipfile.ZipFile(result.workbook_path) as archive:
                embedded = [
                    archive.read(name)
                    for name in archive.namelist()
                    if name.startswith("xl/media/")
                ]
            self.assertGreaterEqual(len(embedded), 3)
            self.assertIn(result.qr_path.read_bytes(), embedded)

    def test_generate_ip_weekly_and_monthly(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = IpTemplateGenerator().generate(MonthlyTemplateRequest(
                year=2026,
                month=7,
                output_dir=Path(tmp),
                projected_gc=(450, 450, 750, 1100, 1050),
                real_gc=(504, 504, 541, 1121, 1110),
                projected_hours=(69, 69, 115, 169, 161),
                real_hours=(88, 89, 112, 175, 179),
                ip_target=6.5,
                template_id="IP-2026-07-TEST",
            ))
            self.assertTrue(result.workbook_path.exists())
            self.assertTrue(result.qr_path.exists())
            self.assertEqual(result.period, "2026-07")
            self.assertEqual(result.template_type, "ip_semanal_mensual")

            workbook = load_workbook(result.workbook_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["IP semanal y mensual", "Configuración"],
            )
            sheet = workbook["IP semanal y mensual"]
            self.assertEqual(sheet["B1"].value, "IP Objetivo Julio 6,50")
            self.assert_code_hidden_from_main_sheet(
                sheet,
                "IP-2026-07-TEST",
            )
            self.assertGreaterEqual(len(sheet._images), 2)
            qr_images = [
                image for image in sheet._images
                if image.anchor._from.col == 9
            ]
            self.assertEqual(len(qr_images), 1)
            self.assertEqual(EMU_to_pixels(qr_images[0].anchor.ext.cx), 48)
            self.assertEqual(sheet.column_dimensions["J"].width, 7)
            self.assertNotIn("K", sheet.column_dimensions)
            self.assertEqual(sheet["A2"].value, "Fecha")
            self.assertEqual(sheet["A5"].value, "Miércoles 01-07")
            self.assertEqual(sheet["B5"].value, 450)
            self.assertEqual(sheet["C5"].value, 504)
            self.assertEqual(
                sheet["D5"].value,
                '=IF(OR(B5="",C5=""),"",C5-B5)',
            )
            self.assertEqual(sheet["H5"].value, '=IFERROR(B5/E5,"")')
            self.assertEqual(sheet["I5"].value, '=IFERROR(C5/F5,"")')
            self.assertEqual(sheet["A10"].value, "Total")
            self.assertEqual(
                sheet["B10"].value,
                '=IF(COUNT(B3:B9)=0,"",SUM(B3:B9))',
            )
            self.assertEqual(sheet["A47"].value, "TOTAL MES")
            self.assertEqual(
                sheet["B47"].value,
                '=IF(COUNT(B10,B19,B28,B37,B46)=0,"",'
                'SUM(B10,B19,B28,B37,B46))',
            )
            self.assertEqual(sheet["H47"].value, '=IFERROR(B47/E47,"")')
            self.assertEqual(sheet["I47"].value, '=IFERROR(C47/F47,"")')
            self.assertEqual(sheet["A2"].fill.fgColor.rgb, "00FFF200")
            self.assertEqual(sheet["A10"].fill.fgColor.rgb, "00000000")
            self.assertEqual(sheet["A10"].font.color.rgb, "00FFFFFF")
            self.assertTrue(sheet.sheet_view.showGridLines)
            self.assertFalse(bool(sheet.print_options.gridLines))
            self.assertFalse(bool(sheet.print_options.gridLinesSet))
            self.assert_thin_gray_border(
                sheet,
                ("C5", "D5", "F5", "G5"),
            )
            self.assertEqual(self.body_merges(sheet, 2, 47), [])
            self.assertIn(sheet["B3"].value, (None, ""))
            self.assert_thin_gray_border(sheet, ("B3",))
            self.assertEqual(sheet.print_area, "'IP semanal y mensual'!$A$1:$J$47")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertGreater(len(sheet.conditional_formatting), 0)
            self.assertEqual(
                workbook["Configuración"]["B8"].value,
                "MCEXCEL|TIPO=IP|CODIGO=IP-2026-07-TEST|PERIODO=2026-07",
            )

            formulas = [
                cell.value
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertFalse(any(
                error in formula
                for formula in formulas
                for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
            ))
            with zipfile.ZipFile(result.workbook_path) as archive:
                embedded = [
                    archive.read(name)
                    for name in archive.namelist()
                    if name.startswith("xl/media/")
                ]
            self.assertGreaterEqual(len(embedded), 3)
            self.assertIn(result.qr_path.read_bytes(), embedded)

    def test_ip_generator_is_registered(self):
        self.assertIn("ip_semanal_mensual", TemplateRegistry().available())


if __name__ == "__main__":
    unittest.main()
