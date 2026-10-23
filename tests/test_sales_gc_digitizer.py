import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from bot_excel.core.digitization_models import SalesGcPhotoReading
from bot_excel.core.models import MonthlyTemplateRequest
from bot_excel.services.sales_gc_digitizer import SalesGcDigitizer
from bot_excel.templates.ventas_gc.generator import SalesGcTemplateGenerator


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class SalesGcDigitizerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.output_dir = Path(self.temp.name)
        self.generated = SalesGcTemplateGenerator().generate(
            MonthlyTemplateRequest(
                year=2026,
                month=6,
                output_dir=self.output_dir,
                template_id="VT-2026-06-ABCD",
            )
        )
        self.reading = SalesGcPhotoReading.model_validate({
            "template_type": "ventas_gc",
            "template_code": "VT-2026-06-ABCD",
            "year": 2026,
            "month": 6,
            "rows": [
                {
                    "day": 1,
                    "real_sales": 345,
                    "real_gc": 793,
                    "sales_confidence": 0.99,
                    "gc_confidence": 0.99,
                },
                {
                    "day": 17,
                    "real_sales": 890,
                    "real_gc": 543,
                    "sales_confidence": 0.99,
                    "gc_confidence": 0.99,
                },
            ],
            "image_quality": "good",
        })

    def tearDown(self):
        self.temp.cleanup()

    def test_writes_sales_and_gc_to_structural_day_rows(self):
        result = SalesGcDigitizer().create_digitized_copy(
            original_path=self.generated.workbook_path,
            reading=self.reading,
            year=2026,
            month=6,
        )
        workbook = load_workbook(result.workbook_path, data_only=False)
        sheet = workbook["CONSOLIDADO"]
        self.assertEqual(sheet["D4"].value, 345)
        self.assertEqual(sheet["H4"].value, 793)
        self.assertEqual(sheet["D21"].value, 890)
        self.assertEqual(sheet["H21"].value, 543)
        workbook.close()

    def test_preserves_formulas_and_calculation_mode(self):
        result = SalesGcDigitizer().create_digitized_copy(
            original_path=self.generated.workbook_path,
            reading=self.reading,
            year=2026,
            month=6,
        )
        workbook = load_workbook(result.workbook_path, data_only=False)
        sheet = workbook["CONSOLIDADO"]
        self.assertEqual(sheet["E4"].value, '=IF(OR(B4="",D4=""),"",D4-C4)')
        self.assertEqual(sheet["F4"].value, '=IF(E4="","",SUM($E$4:E4))')
        self.assertEqual(workbook.calculation.calcMode, "auto")
        workbook.close()

    def test_preserves_logo_qr_and_configuration(self):
        result = SalesGcDigitizer().create_digitized_copy(
            original_path=self.generated.workbook_path,
            reading=self.reading,
            year=2026,
            month=6,
        )
        workbook = load_workbook(result.workbook_path)
        self.assertGreaterEqual(len(workbook["CONSOLIDADO"]._images), 2)
        self.assertIn("Configuración", workbook.sheetnames)
        self.assertEqual(
            workbook["Configuración"]["B6"].value,
            "VT-2026-06-ABCD",
        )
        workbook.close()

    def test_original_binary_is_not_modified(self):
        before = sha256(self.generated.workbook_path)
        result = SalesGcDigitizer().create_digitized_copy(
            original_path=self.generated.workbook_path,
            reading=self.reading,
            year=2026,
            month=6,
        )
        self.assertNotEqual(result.workbook_path, self.generated.workbook_path)
        self.assertEqual(sha256(self.generated.workbook_path), before)
        self.assertIn("_DIGITALIZADA", result.workbook_path.name)

    def test_none_does_not_erase_existing_cell(self):
        workbook = load_workbook(self.generated.workbook_path)
        workbook["CONSOLIDADO"]["D4"] = 111
        workbook.save(self.generated.workbook_path)
        workbook.close()
        reading = self.reading.model_copy(update={
            "rows": [self.reading.rows[0].model_copy(
                update={"real_sales": None, "real_gc": 500}
            )]
        })
        result = SalesGcDigitizer().create_digitized_copy(
            original_path=self.generated.workbook_path,
            reading=reading,
            year=2026,
            month=6,
        )
        copied = load_workbook(result.workbook_path)
        self.assertEqual(copied["CONSOLIDADO"]["D4"].value, 111)
        self.assertEqual(copied["CONSOLIDADO"]["H4"].value, 500)
        copied.close()


if __name__ == "__main__":
    unittest.main()
