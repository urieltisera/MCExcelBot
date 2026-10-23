from __future__ import annotations

import calendar
from datetime import datetime
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.errors import IgnoredError, IgnoredErrors
from openpyxl.worksheet.worksheet import Worksheet

from bot_excel.branding import (
    add_brand_logo,
    add_qr_panel,
    build_qr_payload,
    create_qr_png,
)
from bot_excel.core.ids import build_template_id
from bot_excel.core.models import GeneratedTemplate, MonthlyTemplateRequest
from bot_excel.templates.base import MonthlyTemplateGenerator

MONTHS_ES = (
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
)
WEEKDAYS_ES = (
    "LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO",
)

WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")
LIGHT_GRAY_FILL = PatternFill(fill_type="solid", fgColor="D9D9D9")
MID_GRAY_FILL = PatternFill(fill_type="solid", fgColor="B7B7B7")
DARK_FILL = PatternFill(fill_type="solid", fgColor="000000")


class SalesGcTemplateGenerator(MonthlyTemplateGenerator):
    template_key = "ventas_gc"
    display_name = "Ventas y GC"
    prefix = "VT"

    def generate(self, request: MonthlyTemplateRequest) -> GeneratedTemplate:
        self._validate_request(request)
        days_in_month = calendar.monthrange(request.year, request.month)[1]
        sales = self._normalize_values(request.projected_sales, days_in_month)
        gc = self._normalize_values(request.projected_gc, days_in_month)

        template_id = request.template_id or build_template_id(
            self.prefix, request.year, request.month
        )
        qr_payload = build_qr_payload(
            "VENTAS_GC",
            template_id,
            request.year,
            request.month,
        )

        output_dir = Path(request.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        qr_path = output_dir / f"{template_id}_qr.png"
        create_qr_png(qr_payload, qr_path)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CONSOLIDADO"
        config = workbook.create_sheet("Configuración")

        self._build_config_sheet(
            config, request, template_id, qr_payload, qr_path, sales, gc
        )
        self._build_main_sheet(
            sheet=sheet,
            year=request.year,
            month=request.month,
            days_in_month=days_in_month,
            template_id=template_id,
            qr_path=qr_path,
            sales=sales,
            gc=gc,
        )

        filename = (
            f"Ventas_GC_{MONTHS_ES[request.month - 1].title()}_"
            f"{request.year}_{template_id}.xlsx"
        )
        workbook_path = output_dir / filename
        workbook.save(workbook_path)

        return GeneratedTemplate(
            template_id=template_id,
            workbook_path=workbook_path,
            qr_path=qr_path,
            period=f"{request.year}-{request.month:02d}",
            template_type=self.template_key,
        )

    @staticmethod
    def _validate_request(request: MonthlyTemplateRequest) -> None:
        if not 1 <= request.month <= 12:
            raise ValueError("El mes debe estar entre 1 y 12.")
        if not 2000 <= request.year <= 2100:
            raise ValueError("El año debe estar entre 2000 y 2100.")
        for name, values in (
            ("ventas proyectadas", request.projected_sales),
            ("GC proyectados", request.projected_gc),
        ):
            if len(values) > 31:
                raise ValueError(f"{name}: no puede haber más de 31 valores.")
            if any(value < 0 for value in values):
                raise ValueError(f"{name}: no se permiten valores negativos.")

    @staticmethod
    def _normalize_values(values: Sequence[int | float], days: int) -> list[int | float]:
        normalized = list(values[:days])
        normalized.extend([0] * (days - len(normalized)))
        return normalized

    @staticmethod
    def _side(color: str = "#666666", style: str = "thin") -> Side:
        return Side(style=style, color=color.lstrip("#"))

    @classmethod
    def _border(cls, color: str = "#666666", style: str = "thin") -> Border:
        side = cls._side(color, style)
        return Border(top=side, bottom=side, left=side, right=side)

    @staticmethod
    def _cells(sheet: Worksheet, cell_range: str):
        for row in sheet[cell_range]:
            yield from row

    @classmethod
    def _apply_outline(
        cls,
        sheet: Worksheet,
        cell_range: str,
        color: str = "#000000",
        style: str = "medium",
    ) -> None:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        side = cls._side(color, style)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=column)
                cell.border = Border(
                    left=side if column == min_col else cell.border.left,
                    right=side if column == max_col else cell.border.right,
                    top=side if row == min_row else cell.border.top,
                    bottom=side if row == max_row else cell.border.bottom,
                )

    def _build_config_sheet(
        self,
        config: Worksheet,
        request: MonthlyTemplateRequest,
        template_id: str,
        qr_payload: str,
        qr_path: Path,
        sales: list[int | float],
        gc: list[int | float],
    ) -> None:
        config.merge_cells("A1:E1")
        config["A1"] = "CONFIGURACIÓN — VENTAS Y GC"
        for cell in self._cells(config, "A1:E1"):
            cell.fill = WHITE_FILL
            cell.font = Font(bold=True, color="000000", size=14)
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(bottom=self._side("#000000"))

        metadata = [
            ("Campo", "Valor"),
            ("Año", request.year),
            ("Mes", request.month),
            ("Código", template_id),
            ("Contenido QR", qr_payload),
            ("Tipo", self.display_name),
        ]
        for row_index, values in enumerate(metadata, start=3):
            for column_index, value in enumerate(values, start=1):
                config.cell(row=row_index, column=column_index, value=value)

        config_rows = [("Día", "Venta proyectada", "GC proyectados")]
        config_rows.extend(
            (
                day,
                sales[day - 1] if day <= len(sales) else "",
                gc[day - 1] if day <= len(gc) else "",
            )
            for day in range(1, 32)
        )
        for row_index, values in enumerate(config_rows, start=10):
            for column_index, value in enumerate(values, start=1):
                config.cell(row=row_index, column=column_index, value=value)

        for cell in config[3]:
            if cell.column <= 2:
                cell.font = Font(bold=True)
        for cell in config[10]:
            if cell.column <= 3:
                cell.font = Font(bold=True)
        for cell in self._cells(config, "A3:C41"):
            cell.border = self._border()
        for cell in self._cells(config, "B11:B41"):
            cell.number_format = '"$" #,##0'

        config.column_dimensions["A"].width = 18
        config.column_dimensions["B"].width = 22
        config.column_dimensions["C"].width = 18
        config.column_dimensions["E"].width = 12
        try:
            image = Image(str(qr_path))
            image.width = 48
            image.height = 48
            config.add_image(image, "E3")
        except Exception:
            config["E3"] = qr_payload
            config["E3"].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        config.sheet_view.showGridLines = False
        config.sheet_properties.pageSetUpPr.fitToPage = True
        config.page_setup.orientation = "portrait"
        config.page_setup.fitToWidth = 1
        config.page_setup.fitToHeight = 1
        config.page_margins.left = 0.25
        config.page_margins.right = 0.25
        config.page_margins.top = 0.4
        config.page_margins.bottom = 0.4
        config.print_options.horizontalCentered = True
        config.print_area = "A1:C41"

    def _build_main_sheet(
        self,
        sheet: Worksheet,
        year: int,
        month: int,
        days_in_month: int,
        template_id: str,
        qr_path: Path,
        sales: list[int | float],
        gc: list[int | float],
    ) -> None:
        sheet.merge_cells("B1:J1")
        sheet["B1"] = f"PROYECCION DE VENTAS {MONTHS_ES[month - 1]} {year}"
        for cell in self._cells(sheet, "A1:J1"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(name="Arial", bold=True, color="000000", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=self._side("#000000", "thin"))
        add_brand_logo(sheet, anchor="A1", width=28, height=28)
        add_qr_panel(
            sheet,
            qr_path,
            image_anchor="K1",
        )
        self._apply_outline(sheet, "A1:J1")
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 6

        headers = [
            "DÍA", "FECHA", "VENTA PROY.", "VENTA REAL", "DIFERENCIA",
            "ACUMULADO", "GC'S PROYEC.", "GC'S REAL", "DIFERENCIA", "ACUMULADO",
        ]
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row=3, column=column, value=value)
            cell.fill = DARK_FILL
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=5.5)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=False
            )
            cell.border = self._border("#000000")
        sheet.row_dimensions[3].height = 15

        for day in range(1, 32):
            row = day + 3 if day <= 15 else day + 4
            if day <= days_in_month:
                values = [
                    WEEKDAYS_ES[datetime(year, month, day).weekday()],
                    day,
                    sales[day - 1],
                    "",
                    None,
                    None,
                    gc[day - 1],
                    "",
                    None,
                    None,
                ]
            else:
                values = [""] * 10
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.fill = WHITE_FILL
                cell.font = Font(name="Arial", size=6.5, color="000000")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#808080")

            sheet[f"E{row}"] = f'=IF(OR(B{row}="",D{row}=""),"",D{row}-C{row})'
            sheet[f"I{row}"] = f'=IF(OR(B{row}="",H{row}=""),"",H{row}-G{row})'
            if day <= 15:
                sales_accumulation = f"SUM($E$4:E{row})"
                gc_accumulation = f"SUM($I$4:I{row})"
            else:
                sales_accumulation = f"SUM($E$4:$E$18,$E$20:E{row})"
                gc_accumulation = f"SUM($I$4:$I$18,$I$20:I{row})"
            sheet[f"F{row}"] = (
                f'=IF(E{row}="","",{sales_accumulation})'
            )
            sheet[f"J{row}"] = (
                f'=IF(I{row}="","",{gc_accumulation})'
            )
            if day <= days_in_month:
                sheet[f"A{row}"].fill = DARK_FILL
                sheet[f"A{row}"].font = Font(
                    name="Arial", bold=True, size=6, color="FFFFFF"
                )
            sheet.row_dimensions[row].height = 13

        sheet.merge_cells("A19:B19")
        sheet["A19"] = "1° QUINCENA"
        first_half = {
            "C19": "=SUM(C4:C18)",
            "D19": '=IF(COUNT(D4:D18)=0,"",SUM(D4:D18))',
            "E19": '=IF(D19="","",D19-C19)',
            "F19": '=IFERROR(LOOKUP(2,1/(F4:F18<>""),F4:F18),"")',
            "G19": "=SUM(G4:G18)",
            "H19": '=IF(COUNT(H4:H18)=0,"",SUM(H4:H18))',
            "I19": '=IF(H19="","",H19-G19)',
            "J19": '=IFERROR(LOOKUP(2,1/(J4:J18<>""),J4:J18),"")',
        }
        for coordinate, formula in first_half.items():
            sheet[coordinate] = formula
        for cell in self._cells(sheet, "A19:J19"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(name="Arial", bold=True, size=7, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border("#666666")
        sheet.row_dimensions[19].height = 13

        sheet.merge_cells("A36:B36")
        sheet["A36"] = "TOTAL"
        totals = {
            "C36": "=SUM(C4:C18,C20:C35)",
            "D36": '=IF(COUNT(D4:D18,D20:D35)=0,"",SUM(D4:D18,D20:D35))',
            "E36": '=IF(D36="","",D36-C36)',
            "F36": '=IFERROR(LOOKUP(2,1/(F4:F35<>""),F4:F35),"")',
            "G36": "=SUM(G4:G18,G20:G35)",
            "H36": '=IF(COUNT(H4:H18,H20:H35)=0,"",SUM(H4:H18,H20:H35))',
            "I36": '=IF(H36="","",H36-G36)',
            "J36": '=IFERROR(LOOKUP(2,1/(J4:J35<>""),J4:J35),"")',
        }
        for coordinate, formula in totals.items():
            sheet[coordinate] = formula
        for cell in self._cells(sheet, "A36:J36"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(name="Arial", bold=True, size=7.5, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border("#000000")
        for cell in self._cells(sheet, "A36:B36"):
            cell.fill = DARK_FILL
            cell.font = Font(name="Arial", bold=True, size=7.5, color="FFFFFF")
        sheet.row_dimensions[36].height = 15

        for cell in self._cells(sheet, "C4:F36"):
            cell.number_format = '"$" #,##0;[Red]-"$" #,##0'
        for cell in self._cells(sheet, "G4:J36"):
            cell.number_format = '0;[Red]-0'

        self._apply_outline(sheet, "A3:J36")

        for column, width in {
            "A": 9.5,
            "B": 6,
            "C": 10.5,
            "D": 10,
            "E": 10,
            "F": 10.5,
            "G": 11,
            "H": 9.5,
            "I": 10,
            "J": 10.5,
            "K": 7,
        }.items():
            sheet.column_dimensions[column].width = width

        sheet.freeze_panes = "A4"
        sheet.sheet_view.showGridLines = True
        sheet.sheet_view.zoomScale = 80
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.page_margins.left = 0.2
        sheet.page_margins.right = 0.2
        sheet.page_margins.top = 0.3
        sheet.page_margins.bottom = 0.3
        sheet.print_options.horizontalCentered = True
        sheet.print_area = "A1:K36"
        sheet.ignored_errors = IgnoredErrors(ignoredError=[
            IgnoredError(
                sqref="C19:J19 C36:J36",
                formula=True,
                emptyCellReference=True,
            )
        ])
