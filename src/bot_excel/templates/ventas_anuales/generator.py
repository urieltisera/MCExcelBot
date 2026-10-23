from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet

from bot_excel.branding import (
    add_brand_logo,
    add_qr_image,
    add_qr_panel,
    build_annual_qr_payload,
    create_qr_png,
)
from bot_excel.core.ids import build_annual_template_id
from bot_excel.core.models import (
    AnnualSalesMonth,
    AnnualSalesRequest,
    GeneratedTemplate,
)
from bot_excel.services.annual_consolidation import normalize_annual_months
from bot_excel.templates.base import AnnualTemplateGenerator
from bot_excel.templates.ventas_gc.generator import MONTHS_ES

WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")
INPUT_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
LIGHT_GRAY_FILL = PatternFill(fill_type="solid", fgColor="D9D9D9")
MID_GRAY_FILL = PatternFill(fill_type="solid", fgColor="B7B7B7")
DARK_FILL = PatternFill(fill_type="solid", fgColor="000000")
GREEN_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
PINK_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFE699")


class AnnualSalesTemplateGenerator(AnnualTemplateGenerator):
    template_key = "ventas_consolidadas_anuales"
    display_name = "Ventas consolidadas anuales"
    prefix = "VA"
    generator_version = "1.0"

    def generate(self, request: AnnualSalesRequest) -> GeneratedTemplate:
        self._validate_request(request)
        months = normalize_annual_months(request.months)
        template_id = request.template_id or build_annual_template_id(
            self.prefix,
            request.year,
        )
        qr_payload = build_annual_qr_payload(
            "VENTAS_ANUALES",
            template_id,
            request.year,
        )

        output_dir = Path(request.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        qr_path = output_dir / f"{template_id}_qr.png"
        create_qr_png(qr_payload, qr_path)

        workbook = Workbook()
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        sheet = workbook.active
        sheet.title = "CONSOLIDADO ANUAL"
        config = workbook.create_sheet("Configuración")

        self._build_main_sheet(
            sheet=sheet,
            request=request,
            months=months,
            qr_path=qr_path,
        )
        self._build_config_sheet(
            config=config,
            request=request,
            months=months,
            template_id=template_id,
            qr_payload=qr_payload,
            qr_path=qr_path,
        )

        filename = (
            f"Ventas_Consolidadas_Anuales_{request.year}_{template_id}.xlsx"
        )
        workbook_path = output_dir / filename
        workbook.save(workbook_path)

        return GeneratedTemplate(
            template_id=template_id,
            workbook_path=workbook_path,
            qr_path=qr_path,
            period=str(request.year),
            template_type=self.template_key,
        )

    @staticmethod
    def _validate_request(request: AnnualSalesRequest) -> None:
        if not 2000 <= request.year <= 2100:
            raise ValueError("El año debe estar entre 2000 y 2100.")
        normalize_annual_months(request.months)

    @staticmethod
    def _side(color: str = "#808080", style: str = "thin") -> Side:
        return Side(style=style, color=color.lstrip("#"))

    @classmethod
    def _border(cls, color: str = "#808080", style: str = "thin") -> Border:
        side = cls._side(color, style)
        return Border(top=side, bottom=side, left=side, right=side)

    @staticmethod
    def _cells(sheet: Worksheet, cell_range: str):
        for row in sheet[cell_range]:
            yield from row

    @classmethod
    def _apply_outline(
        cls,
        sheet: Worksheet,
        cell_range: str,
        color: str = "#000000",
        style: str = "medium",
    ) -> None:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        side = cls._side(color, style)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=column)
                cell.border = Border(
                    left=side if column == min_col else cell.border.left,
                    right=side if column == max_col else cell.border.right,
                    top=side if row == min_row else cell.border.top,
                    bottom=side if row == max_row else cell.border.bottom,
                )

    @staticmethod
    def _excel_value(value: int | float | None) -> int | float | str:
        return "" if value is None else value

    def _build_main_sheet(
        self,
        sheet: Worksheet,
        request: AnnualSalesRequest,
        months: tuple[AnnualSalesMonth, ...],
        qr_path: Path,
    ) -> None:
        sheet.merge_cells("B1:L1")
        sheet["B1"] = f"VENTAS CONSOLIDADAS ANUALES {request.year}"
        for cell in self._cells(sheet, "A1:M1"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(name="Arial", bold=True, color="000000", size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=self._side("#000000", "thin"))
        add_brand_logo(sheet, anchor="A1", width=30, height=30)
        add_qr_panel(sheet, qr_path, image_anchor="M1", width=48, height=48)
        self._apply_outline(sheet, "A1:M1")
        sheet.row_dimensions[1].height = 38
        sheet.row_dimensions[2].height = 6

        kpis = (
            ("A3", "VENTA PROYECTADA ANUAL", "B3:C3", "=B21", "currency"),
            ("D3", "VENTA REAL ANUAL", "E3:F3", "=C21", "currency"),
            ("G3", "DIFERENCIA ANUAL", "H3:I3", "=D21", "currency"),
            ("J3", "CUMPLIMIENTO ANUAL", "K3:M3", "=E21", "percent"),
            (
                "A4",
                "MEJOR MES POR VENTA REAL",
                "B4:C4",
                '=IF(COUNT(C9:C20)=0,"",INDEX(A9:A20,MATCH(MAX(C9:C20),C9:C20,0)))',
                "text",
            ),
            (
                "D4",
                "MES CON MAYOR CUMPLIMIENTO",
                "E4:F4",
                '=IF(COUNT(E9:E20)=0,"",INDEX(A9:A20,MATCH(MAX(E9:E20),E9:E20,0)))',
                "text",
            ),
            ("G4", "GC PROYECTADOS ANUALES", "H4:I4", "=G21", "integer"),
            ("J4", "GC REALES ANUALES", "K4:M4", "=H21", "integer"),
            ("A5", "DIFERENCIA ANUAL DE GC", "B5:C5", "=I21", "integer"),
            (
                "D5",
                "PROMEDIO MENSUAL VENTA REAL",
                "E5:F5",
                '=IF(COUNT(C9:C20)=0,"",AVERAGE(C9:C20))',
                "currency",
            ),
        )
        for label_cell, label, value_range, formula, value_kind in kpis:
            start, end = value_range.split(":")
            sheet.merge_cells(value_range)
            label_target = sheet[label_cell]
            label_target.value = label
            label_target.fill = DARK_FILL
            label_target.font = Font(
                name="Arial", bold=True, color="FFFFFF", size=6.5
            )
            label_target.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            label_target.border = self._border("#000000")

            value_target = sheet[start]
            value_target.value = formula
            for cell in self._cells(sheet, value_range):
                cell.fill = LIGHT_GRAY_FILL
                cell.font = Font(
                    name="Arial", bold=True, color="000000", size=8
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = self._border("#808080")
            self._apply_outline(sheet, value_range, style="thin")
            if value_kind == "currency":
                value_target.number_format = '"$" #,##0;[Red]-"$" #,##0'
            elif value_kind == "percent":
                value_target.number_format = "0.0%"
            elif value_kind == "integer":
                value_target.number_format = '0;[Red]-0'

        for row in range(3, 6):
            sheet.row_dimensions[row].height = 26
        sheet.row_dimensions[6].height = 6
        sheet.row_dimensions[7].height = 6

        headers = (
            "MES",
            "VENTA PROYECTADA",
            "VENTA REAL",
            "DIFERENCIA",
            "CUMPLIMIENTO %",
            "VAR. VS MES ANTERIOR %",
            "GC PROYECTADOS",
            "GC REALES",
            "DIFERENCIA GC",
            "TICKET PROMEDIO",
            "ACUM. PROYECTADO",
            "ACUM. REAL",
            "CUMPLIMIENTO ACUM. %",
        )
        for column, label in enumerate(headers, start=1):
            cell = sheet.cell(row=8, column=column, value=label)
            cell.fill = DARK_FILL
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=7)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = self._border("#000000")
        sheet.row_dimensions[8].height = 32

        for month_data in months:
            row = month_data.month + 8
            values = (
                MONTHS_ES[month_data.month - 1],
                self._excel_value(month_data.projected_sales),
                self._excel_value(month_data.real_sales),
                None,
                None,
                None,
                self._excel_value(month_data.projected_gc),
                self._excel_value(month_data.real_gc),
                None,
                self._excel_value(month_data.average_ticket),
                None,
                None,
                None,
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.fill = (
                    INPUT_FILL if column in (2, 3, 7, 8, 10) else WHITE_FILL
                )
                cell.font = Font(name="Arial", color="000000", size=8)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#808080")

            sheet[f"D{row}"] = (
                f'=IF(OR(B{row}="",C{row}=""),"",C{row}-B{row})'
            )
            sheet[f"E{row}"] = f'=IFERROR(C{row}/B{row},"")'
            sheet[f"F{row}"] = (
                "" if month_data.month == 1
                else f'=IFERROR(C{row}/C{row - 1}-1,"")'
            )
            sheet[f"I{row}"] = (
                f'=IF(OR(G{row}="",H{row}=""),"",H{row}-G{row})'
            )
            sheet[f"K{row}"] = f"=SUM($B$9:B{row})"
            sheet[f"L{row}"] = f"=SUM($C$9:C{row})"
            sheet[f"M{row}"] = f'=IFERROR(L{row}/K{row},"")'
            sheet.row_dimensions[row].height = 20

        totals = {
            "A21": "TOTAL ANUAL",
            "B21": '=IF(COUNT(B9:B20)=0,"",SUM(B9:B20))',
            "C21": '=IF(COUNT(C9:C20)=0,"",SUM(C9:C20))',
            "D21": '=IF(OR(B21="",C21=""),"",C21-B21)',
            "E21": '=IFERROR(C21/B21,"")',
            "F21": "",
            "G21": '=IF(COUNT(G9:G20)=0,"",SUM(G9:G20))',
            "H21": '=IF(COUNT(H9:H20)=0,"",SUM(H9:H20))',
            "I21": '=IF(OR(G21="",H21=""),"",H21-G21)',
            "J21": "",
            "K21": "=B21",
            "L21": "=C21",
            "M21": '=IFERROR(L21/K21,"")',
        }
        for coordinate, value in totals.items():
            sheet[coordinate] = value
        for cell in self._cells(sheet, "A21:M21"):
            cell.fill = DARK_FILL
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=8)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border("#000000")
        sheet.row_dimensions[21].height = 22

        for cell_range in ("B9:D21", "J9:L21"):
            for cell in self._cells(sheet, cell_range):
                cell.number_format = '"$" #,##0;[Red]-"$" #,##0'
        for cell_range in ("G9:I21",):
            for cell in self._cells(sheet, cell_range):
                cell.number_format = '0;[Red]-0'
        for cell_range in ("E9:F21", "M9:M21"):
            for cell in self._cells(sheet, cell_range):
                cell.number_format = "0.0%"

        sheet.conditional_formatting.add(
            "D9:D20",
            FormulaRule(formula=['AND(D9<>"",D9>0)'], fill=GREEN_FILL),
        )
        sheet.conditional_formatting.add(
            "D9:D20",
            FormulaRule(formula=['AND(D9<>"",D9<0)'], fill=PINK_FILL),
        )
        sheet.conditional_formatting.add(
            "D9:D20",
            FormulaRule(formula=['AND(D9<>"",D9=0)'], fill=YELLOW_FILL),
        )
        sheet.conditional_formatting.add(
            "E9:E20",
            FormulaRule(formula=['AND(E9<>"",E9>=1)'], fill=GREEN_FILL),
        )
        sheet.conditional_formatting.add(
            "E9:E20",
            FormulaRule(formula=['AND(E9<>"",E9<1)'], fill=PINK_FILL),
        )

        self._apply_outline(sheet, "A8:M21")

        sheet["N8"] = "MES"
        sheet["O8"] = "CUMPLIMIENTO MENSUAL"
        sheet["P8"] = "OBJETIVO 100%"
        for row in range(9, 21):
            sheet[f"N{row}"] = f"=A{row}"
            sheet[f"O{row}"] = f'=IF(E{row}="","",E{row})'
            sheet[f"P{row}"] = 1
            for column in ("O", "P"):
                sheet[f"{column}{row}"].number_format = "0%"
        for column in ("N", "O", "P"):
            sheet.column_dimensions[column].hidden = True

        sales_chart = BarChart()
        sales_chart.type = "col"
        sales_chart.grouping = "clustered"
        sales_chart.overlap = 0
        sales_chart.style = 10
        sales_chart.title = "Ventas proyectadas vs reales"
        sales_chart.y_axis.title = "Ventas"
        sales_chart.x_axis.title = "Mes"
        sales_chart.y_axis.numFmt = '"$"#,##0'
        sales_chart.height = 8.2
        sales_chart.width = 14.2
        sales_data = Reference(
            sheet, min_col=2, max_col=3, min_row=8, max_row=20
        )
        month_categories = Reference(
            sheet, min_col=1, min_row=9, max_row=20
        )
        sales_chart.add_data(sales_data, titles_from_data=True)
        sales_chart.set_categories(month_categories)
        sheet.add_chart(sales_chart, "A24")

        compliance_chart = LineChart()
        compliance_chart.style = 13
        compliance_chart.title = "Cumplimiento mensual"
        compliance_chart.y_axis.title = "Cumplimiento"
        compliance_chart.x_axis.title = "Mes"
        compliance_chart.y_axis.numFmt = "0%"
        compliance_chart.height = 8.2
        compliance_chart.width = 14.2
        compliance_chart.visible_cells_only = False
        compliance_data = Reference(
            sheet, min_col=15, max_col=16, min_row=8, max_row=20
        )
        compliance_categories = Reference(
            sheet, min_col=14, min_row=9, max_row=20
        )
        compliance_chart.add_data(compliance_data, titles_from_data=True)
        compliance_chart.set_categories(compliance_categories)
        sheet.add_chart(compliance_chart, "H24")

        for column, width in {
            "A": 13,
            "B": 14,
            "C": 14,
            "D": 12,
            "E": 13,
            "F": 15,
            "G": 13,
            "H": 11,
            "I": 12,
            "J": 13,
            "K": 14,
            "L": 14,
            "M": 15,
        }.items():
            sheet.column_dimensions[column].width = width

        sheet.freeze_panes = "A9"
        sheet.sheet_view.showGridLines = True
        sheet.sheet_view.zoomScale = 75
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.2
        sheet.page_margins.right = 0.2
        sheet.page_margins.top = 0.3
        sheet.page_margins.bottom = 0.3
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:1"
        sheet.print_area = "A1:M41"
        sheet.row_breaks.append(Break(id=22))

    def _build_config_sheet(
        self,
        config: Worksheet,
        request: AnnualSalesRequest,
        months: tuple[AnnualSalesMonth, ...],
        template_id: str,
        qr_payload: str,
        qr_path: Path,
    ) -> None:
        config.merge_cells("B1:F1")
        config["B1"] = "CONFIGURACIÓN — VENTAS CONSOLIDADAS ANUALES"
        for cell in self._cells(config, "A1:F1"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(name="Arial", bold=True, size=12, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=self._side("#000000", "thin"))
        add_brand_logo(config, anchor="A1", width=30, height=30)
        self._apply_outline(config, "A1:F1")
        config.row_dimensions[1].height = 32

        generated_at = datetime.now()
        metadata = (
            ("Campo", "Valor"),
            ("Tipo de plantilla", self.display_name),
            ("Código único", template_id),
            ("Año", request.year),
            ("Fecha y hora de generación", generated_at),
            ("Payload QR", qr_payload),
            ("Versión del generador", self.generator_version),
            ("Origen de datos mensuales", "Ingreso manual"),
            ("Consolidación automática", "Reservado para futura integración"),
        )
        for row_index, values in enumerate(metadata, start=3):
            for column_index, value in enumerate(values, start=1):
                cell = config.cell(row=row_index, column=column_index, value=value)
                cell.border = self._border("#808080")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cell in self._cells(config, "A3:B3"):
            cell.fill = LIGHT_GRAY_FILL
            cell.font = Font(bold=True, size=9)
        config["B7"].number_format = "yyyy-mm-dd hh:mm:ss"

        headers = (
            "Mes",
            "Venta proyectada",
            "Venta real",
            "GC proyectados",
            "GC reales",
            "Ticket promedio",
        )
        for column, label in enumerate(headers, start=1):
            cell = config.cell(row=14, column=column, value=label)
            cell.fill = DARK_FILL
            cell.font = Font(bold=True, color="FFFFFF", size=8)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = self._border("#000000")
        for month_data in months:
            row = month_data.month + 14
            values = (
                MONTHS_ES[month_data.month - 1],
                self._excel_value(month_data.projected_sales),
                self._excel_value(month_data.real_sales),
                self._excel_value(month_data.projected_gc),
                self._excel_value(month_data.real_gc),
                self._excel_value(month_data.average_ticket),
            )
            for column, value in enumerate(values, start=1):
                cell = config.cell(row=row, column=column, value=value)
                cell.border = self._border("#808080")
                cell.alignment = Alignment(horizontal="center")
                if column in (2, 3, 6):
                    cell.number_format = '"$" #,##0;[Red]-"$" #,##0'
                elif column in (4, 5):
                    cell.number_format = '0;[Red]-0'

        add_qr_image(config, qr_path, anchor="E3", width=48, height=48)
        for column, width in {
            "A": 27,
            "B": 34,
            "C": 18,
            "D": 18,
            "E": 18,
            "F": 18,
        }.items():
            config.column_dimensions[column].width = width
        config.sheet_view.showGridLines = False
        config.sheet_properties.pageSetUpPr.fitToPage = True
        config.page_setup.orientation = "landscape"
        config.page_setup.fitToWidth = 1
        config.page_setup.fitToHeight = 1
        config.page_margins.left = 0.3
        config.page_margins.right = 0.3
        config.print_options.horizontalCentered = True
        config.print_area = "A1:F26"
