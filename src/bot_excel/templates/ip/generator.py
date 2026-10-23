from __future__ import annotations

import calendar
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from bot_excel.branding import (
    add_brand_logo,
    add_qr_image,
    add_qr_panel,
    build_qr_payload,
    create_qr_png,
)
from bot_excel.core.ids import build_template_id
from bot_excel.core.models import GeneratedTemplate, MonthlyTemplateRequest
from bot_excel.templates.base import MonthlyTemplateGenerator
from bot_excel.templates.ventas_gc.generator import MONTHS_ES, WEEKDAYS_ES

WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF200")
NEUTRAL_FILL = PatternFill(fill_type="solid", fgColor="FFE699")
GREEN_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
PINK_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
DARK_FILL = PatternFill(fill_type="solid", fgColor="000000")


class IpTemplateGenerator(MonthlyTemplateGenerator):
    template_key = "ip_semanal_mensual"
    display_name = "IP semanal y mensual"
    prefix = "IP"

    def generate(self, request: MonthlyTemplateRequest) -> GeneratedTemplate:
        self._validate_request(request)
        days_in_month = calendar.monthrange(request.year, request.month)[1]
        projected_gc = self._normalize_values(request.projected_gc, days_in_month)
        real_gc = self._normalize_values(request.real_gc, days_in_month)
        projected_hours = self._normalize_values(
            request.projected_hours, days_in_month
        )
        real_hours = self._normalize_values(request.real_hours, days_in_month)

        template_id = request.template_id or build_template_id(
            self.prefix, request.year, request.month
        )
        qr_payload = build_qr_payload(
            "IP",
            template_id,
            request.year,
            request.month,
        )

        output_dir = Path(request.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        qr_path = output_dir / f"{template_id}_qr.png"
        create_qr_png(qr_payload, qr_path)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "IP semanal y mensual"
        config = workbook.create_sheet("Configuración")

        self._build_main_sheet(
            sheet=sheet,
            request=request,
            template_id=template_id,
            qr_path=qr_path,
            projected_gc=projected_gc,
            real_gc=real_gc,
            projected_hours=projected_hours,
            real_hours=real_hours,
        )
        self._build_config_sheet(
            config=config,
            request=request,
            template_id=template_id,
            qr_payload=qr_payload,
            qr_path=qr_path,
            projected_gc=projected_gc,
            real_gc=real_gc,
            projected_hours=projected_hours,
            real_hours=real_hours,
        )

        filename = (
            f"IP_Semanal_Mensual_{MONTHS_ES[request.month - 1].title()}_"
            f"{request.year}_{template_id}.xlsx"
        )
        workbook_path = output_dir / filename
        workbook.save(workbook_path)

        return GeneratedTemplate(
            template_id=template_id,
            workbook_path=workbook_path,
            qr_path=qr_path,
            period=f"{request.year}-{request.month:02d}",
            template_type=self.template_key,
        )

    @staticmethod
    def _validate_request(request: MonthlyTemplateRequest) -> None:
        if not 1 <= request.month <= 12:
            raise ValueError("El mes debe estar entre 1 y 12.")
        if not 2000 <= request.year <= 2100:
            raise ValueError("El año debe estar entre 2000 y 2100.")
        if not 0 < request.ip_target <= 100:
            raise ValueError("El objetivo IP debe ser mayor que 0 y menor o igual a 100.")
        for name, values in (
            ("GC proyectados", request.projected_gc),
            ("GC reales", request.real_gc),
            ("horas proyectadas", request.projected_hours),
            ("horas reales", request.real_hours),
        ):
            if len(values) > 31:
                raise ValueError(f"{name}: no puede haber más de 31 valores.")
            if any(value < 0 for value in values):
                raise ValueError(f"{name}: no se permiten valores negativos.")

    @staticmethod
    def _normalize_values(
        values: Sequence[int | float],
        days: int,
    ) -> list[int | float | str]:
        normalized: list[int | float | str] = list(values[:days])
        normalized.extend([""] * (days - len(normalized)))
        return normalized

    @staticmethod
    def _side(color: str = "#666666", style: str = "thin") -> Side:
        return Side(style=style, color=color.lstrip("#"))

    @classmethod
    def _border(cls, color: str = "#666666", style: str = "thin") -> Border:
        side = cls._side(color, style)
        return Border(top=side, bottom=side, left=side, right=side)

    @staticmethod
    def _cells(sheet: Worksheet, cell_range: str):
        for row in sheet[cell_range]:
            yield from row

    @classmethod
    def _apply_outline(
        cls,
        sheet: Worksheet,
        cell_range: str,
        color: str = "#000000",
        style: str = "medium",
    ) -> None:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        side = cls._side(color, style)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=column)
                cell.border = Border(
                    left=side if column == min_col else cell.border.left,
                    right=side if column == max_col else cell.border.right,
                    top=side if row == min_row else cell.border.top,
                    bottom=side if row == max_row else cell.border.bottom,
                )

    @staticmethod
    def _format_target(value: float) -> str:
        return f"{value:.2f}".replace(".", ",")

    def _build_main_sheet(
        self,
        sheet: Worksheet,
        request: MonthlyTemplateRequest,
        template_id: str,
        qr_path: Path,
        projected_gc: list[int | float | str],
        real_gc: list[int | float | str],
        projected_hours: list[int | float | str],
        real_hours: list[int | float | str],
    ) -> None:
        sheet.merge_cells("B1:I1")
        sheet["B1"] = (
            f"IP Objetivo {MONTHS_ES[request.month - 1].title()} "
            f"{self._format_target(request.ip_target)}"
        )
        for cell in self._cells(sheet, "A1:I1"):
            cell.fill = WHITE_FILL
            cell.font = Font(name="Arial", bold=True, size=12, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border("#000000")
        add_brand_logo(sheet, anchor="A1", width=30, height=30)
        add_qr_panel(
            sheet,
            qr_path,
            image_anchor="J1",
        )
        self._apply_outline(sheet, "A1:I1")
        sheet.row_dimensions[1].height = 30

        headers = (
            "Fecha",
            "Gc´s Vemi",
            "Gc´s Reales",
            "Diferencia",
            "Horas Vemi",
            "Horas Reales",
            "Diferencia",
            "IP Vemi",
            "IP Real",
        )
        weeks = calendar.Calendar(firstweekday=calendar.MONDAY).monthdayscalendar(
            request.year, request.month
        )
        weekly_total_rows: list[int] = []
        next_row = 2

        for week in weeks:
            header_row = next_row
            for column, label in enumerate(headers, start=1):
                cell = sheet.cell(header_row, column, label)
                cell.fill = YELLOW_FILL
                cell.font = Font(name="Arial", bold=True, size=9, color="000000")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#000000")
            sheet.row_dimensions[header_row].height = 18

            first_daily_row = header_row + 1
            for offset, day in enumerate(week):
                row = first_daily_row + offset
                if day:
                    day_index = day - 1
                    sheet[f"A{row}"] = (
                        f"{WEEKDAYS_ES[offset].title()} "
                        f"{day:02d}-{request.month:02d}"
                    )
                    sheet[f"B{row}"] = projected_gc[day_index]
                    sheet[f"C{row}"] = real_gc[day_index]
                    sheet[f"E{row}"] = projected_hours[day_index]
                    sheet[f"F{row}"] = real_hours[day_index]
                else:
                    for column in ("A", "B", "C", "E", "F"):
                        sheet[f"{column}{row}"] = ""

                sheet[f"D{row}"] = (
                    f'=IF(OR(B{row}="",C{row}=""),"",C{row}-B{row})'
                )
                sheet[f"G{row}"] = (
                    f'=IF(OR(E{row}="",F{row}=""),"",F{row}-E{row})'
                )
                sheet[f"H{row}"] = f'=IFERROR(B{row}/E{row},"")'
                sheet[f"I{row}"] = f'=IFERROR(C{row}/F{row},"")'
                for cell in self._cells(sheet, f"A{row}:I{row}"):
                    cell.fill = WHITE_FILL
                    cell.font = Font(name="Arial", bold=True, size=8, color="000000")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = self._border("#808080")
                sheet.row_dimensions[row].height = 16

            total_row = first_daily_row + 7
            weekly_total_rows.append(total_row)
            sheet[f"A{total_row}"] = "Total"
            for column in ("B", "C", "E", "F"):
                sheet[f"{column}{total_row}"] = (
                    f'=IF(COUNT({column}{first_daily_row}:'
                    f'{column}{total_row - 1})=0,"",'
                    f"SUM({column}{first_daily_row}:{column}{total_row - 1}))"
                )
            sheet[f"D{total_row}"] = (
                f'=IF(OR(B{total_row}="",C{total_row}=""),"",'
                f"C{total_row}-B{total_row})"
            )
            sheet[f"G{total_row}"] = (
                f'=IF(OR(E{total_row}="",F{total_row}=""),"",'
                f"F{total_row}-E{total_row})"
            )
            sheet[f"H{total_row}"] = (
                f'=IFERROR(B{total_row}/E{total_row},"")'
            )
            sheet[f"I{total_row}"] = (
                f'=IFERROR(C{total_row}/F{total_row},"")'
            )
            for cell in self._cells(sheet, f"A{total_row}:I{total_row}"):
                cell.fill = DARK_FILL
                cell.font = Font(name="Arial", bold=True, size=8, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#000000")
            sheet.row_dimensions[total_row].height = 17
            self._apply_conditional_formats(
                sheet, first_daily_row, total_row - 1
            )
            self._apply_outline(sheet, f"A{header_row}:I{total_row}")
            next_row = total_row + 1

        monthly_row = next_row
        sheet[f"A{monthly_row}"] = "TOTAL MES"
        total_refs = ",".join(str(row) for row in weekly_total_rows)
        for column in ("B", "C", "E", "F"):
            refs = ",".join(f"{column}{row}" for row in weekly_total_rows)
            sheet[f"{column}{monthly_row}"] = (
                f'=IF(COUNT({refs})=0,"",SUM({refs}))'
            )
        sheet[f"D{monthly_row}"] = (
            f'=IF(OR(B{monthly_row}="",C{monthly_row}=""),"",'
            f"C{monthly_row}-B{monthly_row})"
        )
        sheet[f"G{monthly_row}"] = (
            f'=IF(OR(E{monthly_row}="",F{monthly_row}=""),"",'
            f"F{monthly_row}-E{monthly_row})"
        )
        sheet[f"H{monthly_row}"] = (
            f'=IFERROR(B{monthly_row}/E{monthly_row},"")'
        )
        sheet[f"I{monthly_row}"] = (
            f'=IFERROR(C{monthly_row}/F{monthly_row},"")'
        )
        for cell in self._cells(sheet, f"A{monthly_row}:I{monthly_row}"):
            cell.fill = DARK_FILL
            cell.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border("#000000", "medium")
        sheet.row_dimensions[monthly_row].height = 19

        for cell in self._cells(sheet, f"B2:G{monthly_row}"):
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
        for cell in self._cells(sheet, f"H2:I{monthly_row}"):
            cell.number_format = '0.00'

        for column, width in {
            "A": 20,
            "B": 14,
            "C": 14,
            "D": 14,
            "E": 14,
            "F": 14,
            "G": 14,
            "H": 14,
            "I": 14,
            "J": 7,
        }.items():
            sheet.column_dimensions[column].width = width

        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = True
        sheet.sheet_view.zoomScale = 75
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.page_margins.left = 0.2
        sheet.page_margins.right = 0.2
        sheet.page_margins.top = 0.25
        sheet.page_margins.bottom = 0.25
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:J{monthly_row}"

    @staticmethod
    def _apply_conditional_formats(
        sheet: Worksheet,
        first_row: int,
        last_row: int,
    ) -> None:
        gc_range = f"D{first_row}:D{last_row}"
        hours_range = f"G{first_row}:G{last_row}"
        ip_range = f"I{first_row}:I{last_row}"

        sheet.conditional_formatting.add(
            gc_range,
            FormulaRule(
                formula=[f"AND(ISNUMBER(D{first_row}),D{first_row}>0)"],
                fill=GREEN_FILL,
                font=Font(color="006100", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            gc_range,
            FormulaRule(
                formula=[f"AND(ISNUMBER(D{first_row}),D{first_row}<0)"],
                fill=PINK_FILL,
                font=Font(color="9C0006", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            gc_range,
            FormulaRule(
                formula=[f"AND(ISNUMBER(D{first_row}),D{first_row}=0)"],
                fill=NEUTRAL_FILL,
                font=Font(color="9C6500", bold=True),
            ),
        )

        sheet.conditional_formatting.add(
            hours_range,
            FormulaRule(
                formula=[f"AND(ISNUMBER(G{first_row}),G{first_row}<0)"],
                fill=GREEN_FILL,
                font=Font(color="006100", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            hours_range,
            FormulaRule(
                formula=[f"AND(ISNUMBER(G{first_row}),G{first_row}>0)"],
                fill=PINK_FILL,
                font=Font(color="9C0006", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            hours_range,
            FormulaRule(
                formula=[f"AND(ISNUMBER(G{first_row}),G{first_row}=0)"],
                fill=NEUTRAL_FILL,
                font=Font(color="9C6500", bold=True),
            ),
        )

        sheet.conditional_formatting.add(
            ip_range,
            FormulaRule(
                formula=[
                    f"AND(ISNUMBER(I{first_row}),ISNUMBER(H{first_row}),"
                    f"I{first_row}>=H{first_row})"
                ],
                fill=GREEN_FILL,
                font=Font(color="006100", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            ip_range,
            FormulaRule(
                formula=[
                    f"AND(ISNUMBER(I{first_row}),ISNUMBER(H{first_row}),"
                    f"I{first_row}<H{first_row})"
                ],
                fill=PINK_FILL,
                font=Font(color="9C0006", bold=True),
            ),
        )

    def _build_config_sheet(
        self,
        config: Worksheet,
        request: MonthlyTemplateRequest,
        template_id: str,
        qr_payload: str,
        qr_path: Path,
        projected_gc: list[int | float | str],
        real_gc: list[int | float | str],
        projected_hours: list[int | float | str],
        real_hours: list[int | float | str],
    ) -> None:
        config.merge_cells("B1:F1")
        config["B1"] = "CONFIGURACIÓN — IP SEMANAL Y MENSUAL"
        for cell in self._cells(config, "A1:F1"):
            cell.fill = WHITE_FILL
            cell.font = Font(name="Arial", bold=True, size=12, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._border("#000000")
        add_brand_logo(config, anchor="A1", width=30, height=30)
        config.row_dimensions[1].height = 32

        metadata = [
            ("Campo", "Valor"),
            ("Año", request.year),
            ("Mes", request.month),
            ("Objetivo IP", request.ip_target),
            ("Código", template_id),
            ("Contenido QR", qr_payload),
            ("Tipo", self.display_name),
        ]
        for row_index, values in enumerate(metadata, start=3):
            for column_index, value in enumerate(values, start=1):
                cell = config.cell(row=row_index, column=column_index, value=value)
                cell.border = self._border("#808080")
        for cell in self._cells(config, "A3:B3"):
            cell.fill = YELLOW_FILL
            cell.font = Font(bold=True, size=9)
        config["B6"].number_format = "0.00"

        headers = (
            "Día",
            "GC proyectados",
            "GC reales",
            "Horas proyectadas",
            "Horas reales",
        )
        for column, label in enumerate(headers, start=1):
            cell = config.cell(11, column, label)
            cell.fill = YELLOW_FILL
            cell.font = Font(bold=True, size=9)
            cell.alignment = Alignment(horizontal="center")
            cell.border = self._border("#000000")
        for day in range(1, 32):
            row = day + 11
            values = (
                day,
                projected_gc[day - 1] if day <= len(projected_gc) else "",
                real_gc[day - 1] if day <= len(real_gc) else "",
                projected_hours[day - 1] if day <= len(projected_hours) else "",
                real_hours[day - 1] if day <= len(real_hours) else "",
            )
            for column, value in enumerate(values, start=1):
                cell = config.cell(row, column, value)
                cell.border = self._border("#808080")
                cell.alignment = Alignment(horizontal="center")

        add_qr_image(config, qr_path, anchor="E3", width=48, height=48)
        for column, width in {
            "A": 18,
            "B": 20,
            "C": 16,
            "D": 20,
            "E": 16,
            "F": 12,
        }.items():
            config.column_dimensions[column].width = width
        config.sheet_view.showGridLines = False
        config.sheet_properties.pageSetUpPr.fitToPage = True
        config.page_setup.orientation = "portrait"
        config.page_setup.fitToWidth = 1
        config.page_setup.fitToHeight = 1
        config.page_margins.left = 0.3
        config.page_margins.right = 0.3
        config.print_options.horizontalCentered = True
        config.print_area = "A1:F42"
