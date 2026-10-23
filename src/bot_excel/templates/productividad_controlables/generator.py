from __future__ import annotations

import calendar
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from bot_excel.branding import (
    add_brand_logo,
    add_qr_image,
    add_qr_panel,
    build_qr_payload,
    create_qr_png,
)
from bot_excel.core.ids import build_template_id
from bot_excel.core.models import GeneratedTemplate, MonthlyTemplateRequest
from bot_excel.templates.base import MonthlyTemplateGenerator
from bot_excel.templates.ventas_gc.generator import MONTHS_ES, WEEKDAYS_ES

WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF")
LIGHT_GRAY_FILL = PatternFill(fill_type="solid", fgColor="D9D9D9")
MID_GRAY_FILL = PatternFill(fill_type="solid", fgColor="B7B7B7")
DARK_FILL = PatternFill(fill_type="solid", fgColor="000000")


class ProductivityControllablesTemplateGenerator(MonthlyTemplateGenerator):
    template_key = "productividad_controlables"
    display_name = "Productividad y controlables"
    prefix = "PC"

    def generate(self, request: MonthlyTemplateRequest) -> GeneratedTemplate:
        self._validate_request(request)
        days_in_month = calendar.monthrange(request.year, request.month)[1]
        template_id = request.template_id or build_template_id(
            self.prefix, request.year, request.month
        )
        qr_payload = build_qr_payload(
            "PRODUCTIVIDAD_CONTROLABLES",
            template_id,
            request.year,
            request.month,
        )

        output_dir = Path(request.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        qr_path = output_dir / f"{template_id}_qr.png"
        create_qr_png(qr_payload, qr_path)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Productividad y controlables"
        config = workbook.create_sheet("Configuración")

        self._build_main_sheet(
            sheet=sheet,
            year=request.year,
            month=request.month,
            days_in_month=days_in_month,
            template_id=template_id,
            qr_path=qr_path,
        )
        self._build_config_sheet(
            config=config,
            request=request,
            template_id=template_id,
            qr_payload=qr_payload,
            qr_path=qr_path,
        )

        filename = (
            f"Productividad_Controlables_{MONTHS_ES[request.month - 1].title()}_"
            f"{request.year}_{template_id}.xlsx"
        )
        workbook_path = output_dir / filename
        workbook.save(workbook_path)

        return GeneratedTemplate(
            template_id=template_id,
            workbook_path=workbook_path,
            qr_path=qr_path,
            period=f"{request.year}-{request.month:02d}",
            template_type=self.template_key,
        )

    @staticmethod
    def _validate_request(request: MonthlyTemplateRequest) -> None:
        if not 1 <= request.month <= 12:
            raise ValueError("El mes debe estar entre 1 y 12.")
        if not 2000 <= request.year <= 2100:
            raise ValueError("El año debe estar entre 2000 y 2100.")

    @staticmethod
    def _side(color: str = "#666666", style: str = "thin") -> Side:
        return Side(style=style, color=color.lstrip("#"))

    @classmethod
    def _border(cls, color: str = "#666666", style: str = "thin") -> Border:
        side = cls._side(color, style)
        return Border(top=side, bottom=side, left=side, right=side)

    @staticmethod
    def _cells(sheet: Worksheet, cell_range: str):
        for row in sheet[cell_range]:
            yield from row

    @classmethod
    def _apply_outline(
        cls,
        sheet: Worksheet,
        cell_range: str,
        color: str = "#000000",
        style: str = "medium",
    ) -> None:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        side = cls._side(color, style)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=column)
                cell.border = Border(
                    left=side if column == min_col else cell.border.left,
                    right=side if column == max_col else cell.border.right,
                    top=side if row == min_row else cell.border.top,
                    bottom=side if row == max_row else cell.border.bottom,
                )

    def _build_main_sheet(
        self,
        sheet: Worksheet,
        year: int,
        month: int,
        days_in_month: int,
        template_id: str,
        qr_path: Path,
    ) -> None:
        sheet.merge_cells("B1:P1")
        sheet["B1"] = (
            f"PROYECCION DE PRODUCTIVIDAD Y CONTROLABLES "
            f"{MONTHS_ES[month - 1]} {year}"
        )
        for cell in self._cells(sheet, "A1:P1"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(name="Arial", bold=True, size=10, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=self._side("#000000", "thin"))
        add_brand_logo(sheet, anchor="A1", width=28, height=28)
        add_qr_panel(
            sheet,
            qr_path,
            image_anchor="Q1",
        )
        self._apply_outline(sheet, "A1:P1")
        sheet.row_dimensions[1].height = 30

        group_headers = (
            ("A2:B2", "PERÍODO"),
            ("C2:G2", "PRODUCTIVIDAD"),
            ("I2:K2", "CONTROLABLES"),
            ("M2:P2", "CATEGORÍAS CONTROLABLES"),
        )
        for cell_range, label in group_headers:
            sheet.merge_cells(cell_range)
            anchor = cell_range.split(":")[0]
            sheet[anchor] = label
            for cell in self._cells(sheet, cell_range):
                cell.fill = WHITE_FILL
                cell.font = Font(
                    name="Arial", bold=True, size=5.5, color="000000"
                )
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#000000")
            self._apply_outline(sheet, cell_range)
        sheet.row_dimensions[2].height = 13

        headers = {
            "A3": "DÍA",
            "B3": "FECHA",
            "C3": "HS PROY.",
            "D3": "HS REALES",
            "E3": "DIFERENCIA",
            "F3": "ACUMULADA",
            "G3": "PRODUCTIVIDAD",
            "I3": "COMPLETO",
            "J3": "INCOMPLETO",
            "K3": "STAT",
            "M3": "CARNES",
            "N3": "BACON",
            "O3": "MEZCLA",
            "P3": "QUESO",
        }
        for coordinate, label in headers.items():
            cell = sheet[coordinate]
            cell.value = label
            cell.fill = DARK_FILL
            cell.font = Font(
                name="Arial", bold=True, size=5.5, color="FFFFFF"
            )
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=False
            )
            cell.border = self._border("#000000")
        sheet.row_dimensions[3].height = 14

        day_rows: dict[int, int] = {}
        for day in range(1, 32):
            row = day + 3 if day <= 15 else day + 4
            day_rows[day] = row
            if day <= days_in_month:
                sheet[f"A{row}"] = WEEKDAYS_ES[datetime(year, month, day).weekday()]
                sheet[f"B{row}"] = day
                sheet[f"C{row}"] = 0
            else:
                sheet[f"A{row}"] = ""
                sheet[f"B{row}"] = ""
                sheet[f"C{row}"] = ""

            sheet[f"E{row}"] = (
                f'=IF(OR(B{row}="",D{row}=""),"",D{row}-C{row})'
            )
            previous_row = day_rows.get(day - 1)
            if previous_row is None:
                sheet[f"F{row}"] = f'=IF(E{row}="","",E{row})'
            else:
                sheet[f"F{row}"] = (
                    f'=IF(E{row}="","",F{previous_row}+E{row})'
                )
            sheet[f"G{row}"] = (
                f'=IF(OR(C{row}="",D{row}="",D{row}=0),"",C{row}/D{row})'
            )
            sheet[f"K{row}"] = (
                f'=IF(OR(I{row}="",J{row}=""),"",I{row}-J{row})'
            )

            for cell_range in (
                f"A{row}:G{row}",
                f"I{row}:K{row}",
                f"M{row}:P{row}",
            ):
                for cell in self._cells(sheet, cell_range):
                    cell.fill = WHITE_FILL
                    cell.font = Font(name="Arial", size=6, color="000000")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = self._border("#808080")

            if day <= days_in_month:
                sheet[f"A{row}"].fill = DARK_FILL
                sheet[f"A{row}"].font = Font(
                    name="Arial", bold=True, size=6, color="FFFFFF"
                )
            sheet.row_dimensions[row].height = 12.5

        sheet.merge_cells("A19:B19")
        sheet["A19"] = "1° QUINCENA"
        first_half = {
            "C19": "=SUM(C4:C18)",
            "D19": '=IF(COUNT(D4:D18)=0,"",SUM(D4:D18))',
            "E19": '=IF(D19="","",D19-C19)',
            "F19": '=IFERROR(LOOKUP(2,1/(F4:F18<>""),F4:F18),"")',
            "G19": '=IF(OR(D19="",D19=0),"",C19/D19)',
            "I19": '=IF(COUNT(I4:I18)=0,"",SUM(I4:I18))',
            "J19": '=IF(COUNT(J4:J18)=0,"",SUM(J4:J18))',
            "K19": '=IF(OR(I19="",J19=""),"",I19-J19)',
            "M19": '=IF(COUNT(M4:M18)=0,"",SUM(M4:M18))',
            "N19": '=IF(COUNT(N4:N18)=0,"",SUM(N4:N18))',
            "O19": '=IF(COUNT(O4:O18)=0,"",SUM(O4:O18))',
            "P19": '=IF(COUNT(P4:P18)=0,"",SUM(P4:P18))',
        }
        for coordinate, formula in first_half.items():
            sheet[coordinate] = formula
        for cell_range in ("A19:G19", "I19:K19", "M19:P19"):
            for cell in self._cells(sheet, cell_range):
                cell.fill = MID_GRAY_FILL
                cell.font = Font(
                    name="Arial", bold=True, size=6, color="000000"
                )
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#666666")
            self._apply_outline(sheet, cell_range)
        sheet.row_dimensions[19].height = 13

        sheet.merge_cells("A36:B36")
        sheet["A36"] = "TOTAL"
        totals = {
            "C36": "=SUM(C4:C18,C20:C35)",
            "D36": '=IF(COUNT(D4:D18,D20:D35)=0,"",SUM(D4:D18,D20:D35))',
            "E36": '=IF(D36="","",D36-C36)',
            "F36": '=IFERROR(LOOKUP(2,1/(F4:F35<>""),F4:F35),"")',
            "G36": '=IF(OR(D36="",D36=0),"",C36/D36)',
            "I36": '=IF(COUNT(I4:I18,I20:I35)=0,"",SUM(I4:I18,I20:I35))',
            "J36": '=IF(COUNT(J4:J18,J20:J35)=0,"",SUM(J4:J18,J20:J35))',
            "K36": '=IF(OR(I36="",J36=""),"",I36-J36)',
            "M36": '=IF(COUNT(M4:M18,M20:M35)=0,"",SUM(M4:M18,M20:M35))',
            "N36": '=IF(COUNT(N4:N18,N20:N35)=0,"",SUM(N4:N18,N20:N35))',
            "O36": '=IF(COUNT(O4:O18,O20:O35)=0,"",SUM(O4:O18,O20:O35))',
            "P36": '=IF(COUNT(P4:P18,P20:P35)=0,"",SUM(P4:P18,P20:P35))',
        }
        for coordinate, formula in totals.items():
            sheet[coordinate] = formula
        for cell_range in ("A36:G36", "I36:K36", "M36:P36"):
            for cell in self._cells(sheet, cell_range):
                cell.fill = MID_GRAY_FILL
                cell.font = Font(
                    name="Arial", bold=True, size=6, color="000000"
                )
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self._border("#000000")
            self._apply_outline(sheet, cell_range)
        for cell in self._cells(sheet, "A36:B36"):
            cell.fill = DARK_FILL
            cell.font = Font(
                name="Arial", bold=True, size=6, color="FFFFFF"
            )
        sheet.row_dimensions[36].height = 14

        for cell in self._cells(sheet, "C4:F36"):
            cell.number_format = '0.00;[Red]-0.00'
        for cell in self._cells(sheet, "G4:G36"):
            cell.number_format = '0.00'
        for cell_range in ("I4:K36", "M4:P36"):
            for cell in self._cells(sheet, cell_range):
                cell.number_format = '0;[Red]-0'

        for cell_range in ("A3:G36", "I3:K36", "M3:P36"):
            self._apply_outline(sheet, cell_range)

        for column, width in {
            "A": 9.5,
            "B": 6,
            "C": 9,
            "D": 9,
            "E": 10,
            "F": 10.5,
            "G": 13,
            "H": 1.5,
            "I": 10.5,
            "J": 10.5,
            "K": 10.5,
            "L": 1.5,
            "M": 7.5,
            "N": 7.5,
            "O": 7.5,
            "P": 7.5,
            "Q": 7,
        }.items():
            sheet.column_dimensions[column].width = width

        sheet.freeze_panes = "A4"
        sheet.sheet_view.showGridLines = True
        sheet.sheet_view.zoomScale = 75
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.page_margins.left = 0.2
        sheet.page_margins.right = 0.2
        sheet.page_margins.top = 0.25
        sheet.page_margins.bottom = 0.25
        sheet.print_options.horizontalCentered = True
        sheet.print_area = "A1:Q36"

    def _build_config_sheet(
        self,
        config: Worksheet,
        request: MonthlyTemplateRequest,
        template_id: str,
        qr_payload: str,
        qr_path: Path,
    ) -> None:
        config.merge_cells("B1:F1")
        config["B1"] = "CONFIGURACIÓN — PRODUCTIVIDAD Y CONTROLABLES"
        for cell in self._cells(config, "A1:F1"):
            cell.fill = MID_GRAY_FILL
            cell.font = Font(bold=True, size=12, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=self._side("#000000", "medium"))
        add_brand_logo(config, anchor="A1", width=30, height=30)
        self._apply_outline(config, "A1:F1")
        config.row_dimensions[1].height = 32

        metadata = [
            ("Campo", "Valor"),
            ("Año", request.year),
            ("Mes", request.month),
            ("Código", template_id),
            ("Contenido QR", qr_payload),
            ("Tipo", self.display_name),
        ]
        for row_index, values in enumerate(metadata, start=3):
            for column_index, value in enumerate(values, start=1):
                cell = config.cell(row=row_index, column=column_index, value=value)
                cell.border = self._border("#808080")
        for cell in self._cells(config, "A3:B3"):
            cell.fill = LIGHT_GRAY_FILL
            cell.font = Font(bold=True, size=9)
        config.column_dimensions["A"].width = 18
        config.column_dimensions["B"].width = 32
        add_qr_image(config, qr_path, anchor="E3", width=48, height=48)
        config.sheet_view.showGridLines = False
        config.sheet_properties.pageSetUpPr.fitToPage = True
        config.page_setup.orientation = "portrait"
        config.page_setup.fitToWidth = 1
        config.page_setup.fitToHeight = 1
        config.page_margins.left = 0.3
        config.page_margins.right = 0.3
        config.print_options.horizontalCentered = True
        config.print_area = "A1:F8"
