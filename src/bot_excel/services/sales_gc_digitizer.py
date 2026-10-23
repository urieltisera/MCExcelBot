from __future__ import annotations

import calendar
import hashlib
import shutil
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from bot_excel.core.digitization_models import SalesGcPhotoReading


class SalesGcDigitizerError(RuntimeError):
    pass


@dataclass(frozen=True)
class DigitizationResult:
    workbook_path: Path
    values_written: int
    fields_omitted: int


class SalesGcDigitizer:
    sheet_name = "CONSOLIDADO"

    def create_digitized_copy(
        self,
        *,
        original_path: Path,
        reading: SalesGcPhotoReading,
        year: int,
        month: int,
    ) -> DigitizationResult:
        original = Path(original_path).resolve(strict=True)
        original_hash = self._sha256(original)
        days_in_month = calendar.monthrange(year, month)[1]

        try:
            workbook = load_workbook(original)
        except Exception as exc:
            raise SalesGcDigitizerError(
                "El archivo original no se pudo abrir."
            ) from exc

        if self.sheet_name not in workbook.sheetnames:
            raise SalesGcDigitizerError(
                "El archivo original no es una planilla válida de Ventas y GC."
            )

        sheet = workbook[self.sheet_name]
        day_rows = self._locate_day_rows(sheet, days_in_month)
        destination = self._next_destination(original)
        shutil.copy2(original, destination)

        values_written = 0
        fields_omitted = 0
        for row_reading in reading.rows:
            row_number = day_rows.get(row_reading.day)
            if row_number is None:
                destination.unlink(missing_ok=True)
                raise SalesGcDigitizerError(
                    f"No pude localizar el día {row_reading.day} en el Excel."
                )
            if row_reading.real_sales is not None:
                sheet.cell(row=row_number, column=4).value = row_reading.real_sales
                values_written += 1
            else:
                fields_omitted += 1
            if row_reading.real_gc is not None:
                sheet.cell(row=row_number, column=8).value = row_reading.real_gc
                values_written += 1
            else:
                fields_omitted += 1

        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        try:
            workbook.save(destination)
        except Exception as exc:
            destination.unlink(missing_ok=True)
            raise SalesGcDigitizerError(
                "No pude crear la copia digitalizada."
            ) from exc
        finally:
            workbook.close()

        if self._sha256(original) != original_hash:
            destination.unlink(missing_ok=True)
            raise SalesGcDigitizerError(
                "La verificación de seguridad del original falló."
            )

        return DigitizationResult(
            workbook_path=destination,
            values_written=values_written,
            fields_omitted=fields_omitted,
        )

    @staticmethod
    def _locate_day_rows(sheet, days_in_month: int) -> dict[int, int]:
        rows: dict[int, int] = {}
        for row_number in range(1, sheet.max_row + 1):
            value = sheet.cell(row=row_number, column=2).value
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            day = int(value)
            if value != day or not 1 <= day <= days_in_month:
                continue
            if day in rows:
                raise SalesGcDigitizerError(
                    f"El día {day} aparece más de una vez en el Excel."
                )
            rows[day] = row_number
        if set(rows) != set(range(1, days_in_month + 1)):
            raise SalesGcDigitizerError(
                "La estructura de días del Excel original es inválida."
            )
        return rows

    @staticmethod
    def _next_destination(original: Path) -> Path:
        preferred = original.with_name(
            f"{original.stem}_DIGITALIZADA{original.suffix}"
        )
        if not preferred.exists():
            return preferred
        for index in range(2, 1000):
            candidate = original.with_name(
                f"{original.stem}_DIGITALIZADA_{index}{original.suffix}"
            )
            if not candidate.exists():
                return candidate
        raise SalesGcDigitizerError(
            "No pude reservar un nombre para la copia digitalizada."
        )

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()
